#!/usr/bin/env python3
"""
extract_masterfile.py — 把《运动服饰机会复盘 Masterfile.xlsx》整理成 skill 使用的干净数据文件。

Masterfile 更新后重新运行即可：
    python scripts/extract_masterfile.py <Masterfile.xlsx 路径> [输出目录，默认 ../data]

输出：
    monthly_panel.csv        月度面板：公司 × 月份 × 市值/Forward 收入/净利率/PS/PE（2015-03 起）
    quarterly_panel.csv      季度热图面板：市值/Forward PS/Forward 净利率（含月度面板没有的公司）
    windows.csv              时间轴上高亮的机会窗口（开始月、结束月、窗口内市值涨幅）
    beta_periods.csv         时间轴上标注的行业 Beta 时期
    cases.json               案例库：类型标签 + 前后指标 + 人事/产品/品牌/渠道/行业Beta/评论
    declines.csv             大幅下跌的时间段

依赖：openpyxl、pandas。只读取 Masterfile，不修改它。
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")

MONTHLY_SHEET = "时间轴（月度） (output)"
HEAT_SHEET = "公司市值-PS-净利率高峰低谷热图 "
SIZE_SHEET = "按机会大小整理"
TYPE_SHEETS = {
    "自身成长（基本面有好的变化，做正确的事）": "自身成长",
    "困境反转（陷入困境，基本面有好的变化）": "困境反转",
    "行业beta": "行业Beta",
}
METRIC_LABELS = {"收入 Forward": "fwd_rev", "利润率 Forward": "fwd_npm", "PS Forward": "fwd_ps", "PE Forward": "fwd_pe"}

# 同一公司在不同 sheet 里的写法 → 统一名称
ALIASES = {
    "Adidas（Euro）": "Adidas", "LuLulemon": "LULU", "Decker": "DECK", "Ralph Lauren": "RL",
    "Amer": "AS", "AS（Amer Sports）": "AS", "On Runing": "ON", "斯凯奇": "SKX", "吉登运动服": "GII",
    "Levis": "LEVI", "CROX": "CROX", "卡骆驰": "CROX", "Abercrombie & Fitch": "ANF", "Puma": "PUM",
    "KONTOOR": "KTB", "安德玛": "UAA", "哈尼斯": "HBI", "维密": "VSCO", "Wolverine Worldwide": "WWW",
    "361度": "361度", "宝尊电商": "BZUN", "GOLDWIN": "GOLDWIN", "哥伦比亚户外": "COLM",
    "史蒂夫·马登": "SHOO", "ERMENEGILDO ZEGNA": "ZGN", "亚瑟士": "亚瑟士",
}


# 案例键 = "公司|窗口"。Masterfile 里的笔误或不同 sheet 写法不一致，在这里合并。
WINDOW_ALIASES = {
    "ON|24Q4–24Q4": "ON|23Q1–24Q4",
}
# 类型 sheet 未覆盖、或分组被上一行标签误带的案例。依据写在 note 里，更新 Masterfile 后可删改。
CASE_OVERRIDES = {
    "安踏|15Q1–21Q2": {"copy_from": "安踏|18Q4–21Q2",
                      "note": "长窗口版本；分类沿用 18Q4–21Q2"},
    "比音勒芬|21Q1–23Q2": {"types": ["行业Beta"], "groups": ["中国高尔夫市场高增长（21–24 年 CAGR 20–30%）+ 疫情刺激"],
                        "note": "Masterfile 中该行无分组标签，原先被上方 value brand 分组误带"},
    "GII|20Q1–21Q4": {"types": ["行业Beta"], "groups": ["2020 年疫情放水反弹"],
                     "note": "窗口起点为 2020M3，属疫情后放水反弹"},
    "HBI|20Q1–21Q1": {"types": ["行业Beta"], "groups": ["2020 年疫情放水反弹"],
                     "note": "窗口起点为 2020M3，属疫情后放水反弹"},
    "LEVI|20Q3–21Q2": {"types": ["行业Beta"], "groups": ["2020 年疫情放水反弹"],
                      "note": "疫情后放水反弹"},
    "波司登|20Q1–21Q3": {"types": ["行业Beta"], "groups": ["2020 年疫情放水反弹"],
                       "note": "窗口起点为 2020M3，属疫情后放水反弹；品牌升级在 18Q2–19Q3 已兑现"},
    "CROX|17Q2–19Q1": {"types": ["困境反转"], "groups": ["2014 年起改革的底部公司，利润率从负值修复"],
                      "note": "按原文：13–15 年困境、14 年改革、17 年股价开始反弹；Masterfile 类型 sheet 未收录"},
    "WWW|16Q1–18Q3": {"types": ["困境反转"], "groups": ["利润率先修复、收入拐点未至"],
                     "note": "按原文：16–18 年重组后盈利拐点驱动估值，19 年收入拐点没来，股价再跌"},
    "GOLDWIN|17Q1–19Q3": {"types": ["未分类"], "groups": [],
                         "note": "Masterfile 与原文均无定性描述，仅有前后指标"},
}


def canon(name):
    if name is None:
        return None
    name = str(name).strip()
    return ALIASES.get(name, name)


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return None


def month_label(s):
    """'15M3' -> '2015-03'"""
    m = re.match(r"^(\d\d)M(\d{1,2})$", str(s).strip())
    return f"20{m.group(1)}-{int(m.group(2)):02d}" if m else None


def quarter_label(s):
    """'3Q15' -> '2015Q3'"""
    m = re.match(r"^(\d)Q(\d\d)$", str(s).strip())
    return f"20{m.group(2)}Q{m.group(1)}" if m else None


def extract_monthly(wb_v, wb_f):
    ws, wsf = wb_v[MONTHLY_SHEET], wb_f[MONTHLY_SHEET]
    months = {c: month_label(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)}
    months = {c: m for c, m in months.items() if m}
    rows, windows, current = [], [], None
    end_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "基础胜率")
    for r in range(3, end_row):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        label = str(label).strip()
        if label in METRIC_LABELS:
            if current is None:
                continue
            key = METRIC_LABELS[label]
            for c, m in months.items():
                v = num(ws.cell(r, c).value)
                if key == "fwd_rev" and v == 0:
                    v = None
                current["data"].setdefault(m, {})[key] = v
            continue
        if label in ("行业beta",):
            continue
        # 公司行：市值
        if current:
            rows.append(current)
        current = {"company": canon(label), "raw": label, "data": {}}
        for c, m in months.items():
            current["data"].setdefault(m, {})["mkt_cap"] = num(ws.cell(r, c).value)
        # 高亮窗口（主题色填充）
        start = prev = None
        for c in sorted(months):
            f = wsf.cell(r, c).fill
            on = bool(f and f.fill_type)
            if on and start is None:
                start = months[c]
            if not on and start is not None:
                windows.append({"company": canon(label), "start": start, "end": prev})
                start = None
            prev = months[c]
        if start is not None:
            windows.append({"company": canon(label), "start": start, "end": prev})
    if current:
        rows.append(current)

    recs = []
    for co in rows:
        for m, d in co["data"].items():
            recs.append({"company": co["company"], "month": m, **d})
    panel = pd.DataFrame(recs)
    for k in ["mkt_cap", "fwd_rev", "fwd_npm", "fwd_ps", "fwd_pe"]:
        if k not in panel:
            panel[k] = None
    panel = panel[["company", "month", "mkt_cap", "fwd_rev", "fwd_npm", "fwd_ps", "fwd_pe"]]
    panel = panel.dropna(subset=["mkt_cap", "fwd_rev", "fwd_npm", "fwd_ps", "fwd_pe"], how="all")
    panel = panel.sort_values(["company", "month"]).reset_index(drop=True)

    win = pd.DataFrame(windows)
    if not win.empty:
        mc = panel.set_index(["company", "month"])["mkt_cap"]
        win["mkt_cap_start"] = [mc.get((a, s)) for a, s in zip(win.company, win.start)]
        win["mkt_cap_end"] = [mc.get((a, e)) for a, e in zip(win.company, win.end)]
        win["return"] = win.mkt_cap_end / win.mkt_cap_start - 1

    # 行业 Beta 标注（第 2、3 行）
    betas = []
    for r in (2, 3):
        for c, m in months.items():
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                betas.append({"month": m, "description": v.strip()})
    return panel, win, pd.DataFrame(betas)


def extract_quarterly(wb_v):
    ws = wb_v[HEAT_SHEET]
    cols = {c: canon(ws.cell(1, c).value) for c in range(3, 40) if ws.cell(1, c).value}
    blocks = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a in ("市值", "Forward PS", "Forward 净利率"):
            blocks[{"市值": "mkt_cap", "Forward PS": "fwd_ps", "Forward 净利率": "fwd_npm"}[a]] = r
    recs = {}
    for key, r0 in blocks.items():
        r = r0
        while r <= ws.max_row:
            q = quarter_label(ws.cell(r, 2).value)
            if not q:
                break
            for c, co in cols.items():
                recs.setdefault((co, q), {})[key] = num(ws.cell(r, c).value)
            r += 1
    df = pd.DataFrame([{"company": k[0], "quarter": k[1], **v} for k, v in recs.items()])
    return df.dropna(subset=[c for c in ["mkt_cap", "fwd_ps", "fwd_npm"] if c in df], how="all") \
             .sort_values(["company", "quarter"]).reset_index(drop=True)


def txt(v):
    return str(v).strip() if isinstance(v, str) and v.strip() else None


def extract_cases(wb_v):
    cases = {}

    def key(co, window):
        w = re.sub(r"\s", "", str(window)).replace("-", "–")
        return f"{canon(co)}|{w}"

    # 1) 按机会大小整理：前后指标
    ws = wb_v[SIZE_SHEET]
    bucket, declines = None, []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip():
            bucket = a.strip()
        co, win = ws.cell(r, 3).value, ws.cell(r, 4).value
        if not co or co in ("公司/股票代码",) or not win:
            continue
        if bucket == "下跌":
            declines.append({"company": canon(co), "window": str(win).replace("-", "–"),
                             "return": num(ws.cell(r, 5).value)})
            continue
        k = key(co, win)
        cases[k] = {
            "company": canon(co), "window": str(win).replace("-", "–"), "size_bucket": bucket,
            "return_multiple": num(ws.cell(r, 5).value), "cagr": num(ws.cell(r, 6).value),
            "before": {"mkt_cap": num(ws.cell(r, 7).value), "fwd_rev": num(ws.cell(r, 9).value),
                       "fwd_npm": num(ws.cell(r, 10).value), "fwd_pe": num(ws.cell(r, 11).value),
                       "fwd_ps": num(ws.cell(r, 12).value)},
            "after": {"mkt_cap": num(ws.cell(r, 8).value), "fwd_rev": num(ws.cell(r, 13).value),
                      "fwd_npm": num(ws.cell(r, 14).value), "fwd_pe": num(ws.cell(r, 15).value),
                      "fwd_ps": num(ws.cell(r, 16).value)},
            "types": [], "groups": [],
        }

    # 2) 三个类型 sheet：类型标签、分组说明、定性字段
    for sheet, typ in TYPE_SHEETS.items():
        ws = wb_v[sheet]
        hdr = {str(ws.cell(1, c).value).strip(): c for c in range(1, 12) if ws.cell(1, c).value}
        group = None
        for r in range(2, ws.max_row + 1):
            a = txt(ws.cell(r, 1).value)
            if a:
                group = a
            co, win = ws.cell(r, 2).value, ws.cell(r, 3).value
            if not co or not win or not isinstance(co, str):
                continue
            k = key(co, win)
            c = cases.setdefault(k, {"company": canon(co), "window": str(win).replace("-", "–"),
                                     "types": [], "groups": [],
                                     "return_multiple": num(ws.cell(r, 4).value),
                                     "before": {"mkt_cap": num(ws.cell(r, 11).value), "fwd_rev": num(ws.cell(r, 12).value),
                                                "fwd_npm": num(ws.cell(r, 13).value), "fwd_pe": num(ws.cell(r, 14).value),
                                                "fwd_ps": num(ws.cell(r, 15).value)},
                                     "after": {"mkt_cap": num(ws.cell(r, 16).value), "fwd_rev": num(ws.cell(r, 17).value),
                                               "fwd_npm": num(ws.cell(r, 18).value), "fwd_pe": num(ws.cell(r, 19).value),
                                               "fwd_ps": num(ws.cell(r, 20).value)}})
            if typ not in c["types"]:
                c["types"].append(typ)
            if group and group not in c["groups"]:
                c["groups"].append(group)
            for field in ("人事", "产品", "品牌", "渠道", "行业Beta", "评论"):
                col = hdr.get(field)
                v = txt(ws.cell(r, col).value) if col else None
                if v and not c.get(field):
                    c[field] = v
    # 手工修正：Masterfile 中的笔误、跨 sheet 窗口写法不一致、未分类案例
    for bad, good in WINDOW_ALIASES.items():
        if bad in cases:
            src_c = cases.pop(bad)
            dst = cases.setdefault(good, src_c)
            for f in ("types", "groups"):
                dst[f] = list(dict.fromkeys((dst.get(f) or []) + (src_c.get(f) or [])))
            for f in ("人事", "产品", "品牌", "渠道", "行业Beta", "评论"):
                if src_c.get(f) and not dst.get(f):
                    dst[f] = src_c[f]
    for k, ov in CASE_OVERRIDES.items():
        if k in cases:
            c = cases[k]
            if "copy_from" in ov and ov["copy_from"] in cases:
                src_c = cases[ov["copy_from"]]
                for f in ("types", "groups", "人事", "产品", "品牌", "渠道", "行业Beta", "评论"):
                    if src_c.get(f) and not c.get(f):
                        c[f] = src_c[f]
            if "types" in ov:
                c["types"] = ov["types"]
            if "groups" in ov:
                c["groups"] = ov["groups"]
            if "note" in ov:
                c["note"] = ov["note"]
    # 假反转标签
    for c in cases.values():
        if any("假反转" in g for g in c.get("groups", [])) or "假反转" in (c.get("评论") or "")[:10]:
            if "假反转" not in c["types"]:
                c["types"].append("假反转")
    # 收益分解（after / before）
    for c in cases.values():
        b, a = c.get("before", {}), c.get("after", {})
        def ratio(k):
            try:
                return a[k] / b[k] if a.get(k) and b.get(k) and b[k] > 0 else None
            except (TypeError, ZeroDivisionError):
                return None
        c["decomposition"] = {"mkt_cap_x": ratio("mkt_cap"), "rev_x": ratio("fwd_rev"),
                              "npm_x": ratio("fwd_npm"), "ps_x": ratio("fwd_ps"), "pe_x": ratio("fwd_pe")}
    return list(cases.values()), pd.DataFrame(declines)


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if not src or not src.exists():
        sys.exit("用法：python extract_masterfile.py <Masterfile.xlsx> [输出目录]")
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "data"
    out.mkdir(parents=True, exist_ok=True)
    wb_v = openpyxl.load_workbook(src, data_only=True)
    wb_f = openpyxl.load_workbook(src)  # 读取填充色（机会窗口）

    panel, windows, betas = extract_monthly(wb_v, wb_f)
    quarterly = extract_quarterly(wb_v)
    cases, declines = extract_cases(wb_v)

    # 把案例与时间轴上的高亮窗口对上（同公司、时间重叠最多者）
    def q_to_months(w):
        m = re.match(r"^(\d\d)Q(\d)[–-](\d\d)Q(\d)$", w.replace(" ", ""))
        if not m:
            return None, None
        y1, q1, y2, q2 = m.groups()
        return f"20{y1}-{3 * int(q1) - 2:02d}", f"20{y2}-{3 * int(q2):02d}"

    for c in cases:
        s, e = q_to_months(c["window"])
        c["panel_window"] = None
        if not s or windows.empty:
            continue
        best, best_ov = None, 0
        for _, w in windows[windows.company == c["company"]].iterrows():
            ov_start, ov_end = max(s, w.start), min(e, w.end)
            if ov_start <= ov_end:
                ov = (int(ov_end[:4]) * 12 + int(ov_end[5:])) - (int(ov_start[:4]) * 12 + int(ov_start[5:]))
                if ov >= best_ov:
                    best, best_ov = w, ov
        if best is not None:
            c["panel_window"] = {"start": best.start, "end": best.end,
                                 "return": None if pd.isna(best["return"]) else round(float(best["return"]), 3)}

    panel.to_csv(out / "monthly_panel.csv", index=False)
    quarterly.to_csv(out / "quarterly_panel.csv", index=False)
    windows.to_csv(out / "windows.csv", index=False)
    betas.to_csv(out / "beta_periods.csv", index=False)
    declines.to_csv(out / "declines.csv", index=False)
    (out / "cases.json").write_text(json.dumps(cases, ensure_ascii=False, indent=1))

    print(f"monthly_panel: {panel.company.nunique()} 家公司, {panel.month.min()} ~ {panel.month.max()}, {len(panel)} 行")
    print(f"quarterly_panel: {quarterly.company.nunique()} 家公司, {len(quarterly)} 行")
    print(f"windows: {len(windows)} 个高亮窗口; beta_periods: {len(betas)}; cases: {len(cases)}; declines: {len(declines)}")
    print(f"输出目录：{out}")


if __name__ == "__main__":
    main()
