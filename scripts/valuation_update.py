#!/usr/bin/env python3
"""
valuation_update.py — 估值 → Valuation screening.xlsx（FMP 优先，yfinance fallback）

每 ticker 一个 sheet + Summary 汇总 Forward PE/PS。
"""
import argparse, datetime as dt, sys, time
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from fmp_client import FMPClient, FMPError

MM = 1e6

# ── Tickers ──────────────────────────────────────────────────────
SPORTSWEAR = [
    ("NKE.N","NKE"),("ADDY","ADS.DE"),("ZGN.N","ZGN"),("WWW.N","WWW"),
    ("UAA.N","UAA"),("SHOO.O","SHOO"),("RL.N","RL"),("PUM.DF","PUM.DE"),
    ("ONON.N","ONON"),("NWL.O","NWL"),("LULU.O","LULU"),("LEVI.N","LEVI"),
    ("KTB.N","KTB"),("GIL.N","GIL"),("GAP.N","GAP"),("DECK.N","DECK"),
    ("CROX.O","CROX"),("COLM.O","COLM"),("BZUN.O","BZUN"),("AS.N","AS"),
    ("ANF.N","ANF"),("VSCO.N","VSCO"),("9983.T","9983.T"),("8111.T","8111.T"),("7936.T","7936.T"),
    ("3998.HK","3998.HK"),("2331.HK","2331.HK"),("2020.HK","2020.HK"),
    ("1836.HK","1836.HK"),("1368.HK","1368.HK"),("1361.HK","1361.HK"),
    ("002832.SZ","002832.SZ"),
]
LUXURY = [
    ("VFC","VFC"),("CFR","CFR.SW"),("1913.HK","1913.HK"),("MONC","MONC.MI"),
    ("KER","KER.PA"),("SFER","SFER.MI"),("RACE","RACE"),("RMS","RMS.PA"),
    ("MC","MC.PA"),("BRBY","BRBY.L"),("CPRI","CPRI"),("PVH","PVH"),
    ("TPR","TPR"),("GOOS","GOOS.TO"),
]
ALL_TICKERS = SPORTSWEAR + LUXURY
D2Y = {d: y for d, y in ALL_TICKERS}

# ── 已知报表币种（quote_ccy ≠ financial_ccy 的 ticker）──────────
# yfinance fast_info.currency 是报价币种，财报币种可能不同
KNOWN_FIN_CCY = {
    "BRBY.L":   "GBP",     # 报价 GBp（便士），财报 GBP
    "2020.HK":  "CNY",     # 安踏
    "2331.HK":  "CNY",     # 李宁
    "3998.HK":  "CNY",     # 波司登
    "1368.HK":  "CNY",     # 特步
    "1361.HK":  "CNY",     # 361度
    "002832.SZ":"CNY",     # 比音勒芬
    "1913.HK":  "EUR",     # Prada
    "1836.HK":  "USD",     # 九兴控股
    "CFR.SW":   "EUR",     # Richemont
}

_fx_cache = {}
_fmp_client = FMPClient()

def get_fx(from_ccy, to_ccy):
    """返回 1 from_ccy = ? to_ccy（FMP 优先，yfinance fallback）"""
    if from_ccy == to_ccy:
        return 1.0
    if from_ccy == "GBp" and to_ccy == "GBP":
        return 0.01
    if from_ccy == "GBp" and to_ccy != "GBP":
        return 0.01 * get_fx("GBP", to_ccy)
    key = (from_ccy, to_ccy)
    if key in _fx_cache:
        return _fx_cache[key]
    # 1. 尝试 FMP
    rate = _fmp_client.fx_rate(from_ccy, to_ccy)
    if rate is not None and rate != 1.0:
        _fx_cache[key] = rate
        print(f"  FX {from_ccy}→{to_ccy} = {rate:.4f} (FMP)")
        return rate
    # 2. 回退 yfinance
    try:
        import yfinance as yf
        for pair, inv in [(f"{from_ccy}{to_ccy}=X", False), (f"{to_ccy}{from_ccy}=X", True)]:
            try:
                r = yf.Ticker(pair).fast_info.last_price
                if r and r > 0:
                    result = 1.0 / r if inv else r
                    _fx_cache[key] = result
                    print(f"  FX {from_ccy}→{to_ccy} = {result:.4f} (yfinance)")
                    return result
            except Exception:
                continue
    except ImportError:
        pass
    print(f"  ⚠ FX {from_ccy}→{to_ccy} 获取失败，默认 1.0")
    _fx_cache[key] = 1.0
    return 1.0

# Row layout constants
R_FY     = 5
R_REV    = 6
R_YOY    = 7
R_GM     = 8
R_OPM    = 9
R_NI     = 10
R_NPM    = 11
R_EPS    = 12
R_SHR    = 13
R_DATE   = 15
R_PRICE  = 16
R_MCAP   = 17
R_PE     = 18
R_PS     = 19
R_ANLST  = 21
ANN_START = 3   # col C
Q_START   = 10  # col J


def safe(func, retries=3, delay=2):
    for i in range(retries):
        try:
            return func()
        except Exception:
            if i < retries - 1:
                time.sleep(delay * (i + 1))
    return None


def _v(df, row, col):
    if df is None or df.empty or row not in df.index or col not in df.columns:
        return None
    v = df.loc[row, col]
    if isinstance(v, (int, float, np.integer, np.floating)) and not np.isnan(v):
        return float(v)
    return None


def _first(df, *rows, col=None):
    for r in rows:
        v = _v(df, r, col)
        if v is not None:
            return v
    return None


def qlabel(d):
    q = (d.month - 1) // 3 + 1
    return f"{d.year % 100}Q{q}"


def _ratio(a, b):
    return a / b if (a and b and b != 0) else None


# ══════════════════════════════════════════════════════════════════
def _fmp_to_yf_income(rows):
    """将 FMP income-statement list 转为 yfinance 格式的 DataFrame（行=指标, 列=日期）。"""
    if not rows:
        return None
    records = {}
    for r in rows:
        col = pd.Timestamp(r["date"])
        records[col] = {
            "Total Revenue": r.get("revenue"),
            "Gross Profit": r.get("grossProfit"),
            "Operating Income": r.get("operatingIncome"),
            "Net Income": r.get("netIncome"),
            "Diluted EPS": r.get("epsdiluted") or r.get("eps"),
            "Basic EPS": r.get("eps"),
        }
    df = pd.DataFrame(records)
    df = df.reindex(columns=sorted(df.columns))
    return df


def _fmp_to_yf_bs(rows):
    """将 FMP balance-sheet list 转为 yfinance 格式的 DataFrame。"""
    if not rows:
        return None
    records = {}
    for r in rows:
        col = pd.Timestamp(r["date"])
        records[col] = {
            "Ordinary Shares Number": r.get("commonStock") or r.get("ordinarySharesNumber"),
            "Share Issued": r.get("commonStock"),
        }
    df = pd.DataFrame(records)
    df = df.reindex(columns=sorted(df.columns))
    return df


def _fmp_estimates_to_yf(est_list, last_end, shares):
    """将 FMP analyst-estimates 转为 yfinance 兼容的 eps_est/rev_est DataFrames。"""
    if not est_list:
        return None, None
    # 筛选未来期
    fwd = [e for e in est_list
           if e.get("date") and pd.Timestamp(e["date"]).date() > last_end + dt.timedelta(days=45)]
    fwd = sorted(fwd, key=lambda e: e["date"])[:2]
    if not fwd:
        return None, None

    eps_rows = {}
    rev_rows = {}
    for i, e in enumerate(fwd):
        period = "0y" if i == 0 else "+1y"
        eps_avg = e.get("estimatedEpsAvg") or e.get("epsAvg")
        rev_avg = e.get("estimatedRevenueAvg") or e.get("revenueAvg")
        n_rev = e.get("numAnalystsRevenue") or e.get("numberAnalystEstimatedRevenue")
        n_eps = e.get("numAnalystsEps") or e.get("numberAnalystsEstimatedEps")
        eps_rows[period] = {"avg": eps_avg, "numberOfAnalysts": n_eps}
        rev_rows[period] = {"avg": rev_avg, "numberOfAnalysts": n_rev}

    eps_est = pd.DataFrame(eps_rows).T if eps_rows else None
    rev_est = pd.DataFrame(rev_rows).T if rev_rows else None
    return eps_est, rev_est


def fetch_one_fmp(code):
    """通过 FMP 拉取一个 ticker 的全部数据，返回与 yfinance fetch_one 相同结构的 dict。"""
    try:
        inc_raw = _fmp_client.income_statement(code, period="annual", limit=5)
        if not inc_raw:
            return None
        inc = _fmp_to_yf_income(inc_raw)
        if inc is None or inc.empty:
            return None

        q_inc_raw = _fmp_client.income_statement(code, period="quarter", limit=6)
        q_inc = _fmp_to_yf_income(q_inc_raw)

        bs_raw = _fmp_client.balance_sheet(code, period="annual", limit=5)
        bs = _fmp_to_yf_bs(bs_raw)

        q_bs_raw = _fmp_client.balance_sheet(code, period="quarter", limit=6)
        q_bs = _fmp_to_yf_bs(q_bs_raw)

        quote = _fmp_client.quote(code)
        profile = _fmp_client.profile(code)

        d = {"yahoo": code}
        d["inc"] = inc
        d["q_inc"] = q_inc
        d["bs"] = bs
        d["q_bs"] = q_bs

        d["price"] = quote.get("price")
        d["mcap_raw"] = quote.get("marketCap") or profile.get("mktCap")
        d["shares"] = (quote.get("sharesOutstanding")
                       or profile.get("sharesOutstanding")
                       or (d.get("mcap_raw") / d["price"] if d.get("mcap_raw") and d.get("price") and d["price"] > 0 else None))
        d["quote_ccy"] = profile.get("currency", "")
        d["name"] = profile.get("companyName") or code

        # 确定财报币种
        rep_ccy = inc_raw[-1].get("reportedCurrency") if inc_raw else ""
        d["fin_ccy"] = KNOWN_FIN_CCY.get(code, rep_ccy or d["quote_ccy"])
        if not d["fin_ccy"]:
            d["fin_ccy"] = d["quote_ccy"]

        quote_ccy = d["quote_ccy"]
        fin_ccy = d["fin_ccy"]
        if d["mcap_raw"] and quote_ccy and fin_ccy and quote_ccy != fin_ccy:
            fx = get_fx(quote_ccy, fin_ccy)
            d["mcap"] = d["mcap_raw"] * fx
            d["ccy"] = fin_ccy
            print(f"  MCap 换算: {quote_ccy} → {fin_ccy} (×{fx:.4f})")
        else:
            d["mcap"] = d["mcap_raw"]
            d["ccy"] = fin_ccy or quote_ccy

        # 一致预期
        last_end = sorted(inc.columns)[-1].date()
        try:
            est_raw = _fmp_client.analyst_estimates(code, period="annual", limit=10)
        except FMPError:
            est_raw = []
        eps_est, rev_est = _fmp_estimates_to_yf(est_raw, last_end, d.get("shares"))
        d["eps_est"] = eps_est
        d["rev_est"] = rev_est

        return d
    except FMPError as e:
        print(f"  FMP 失败 {code}: {e}")
        return None


def fetch_one_yf(code):
    """yfinance fallback。"""
    try:
        import yfinance as yf
    except ImportError:
        print(f"  yfinance 不可用，跳过 {code}")
        return None
    t = yf.Ticker(code)
    d = {"yahoo": code}
    d["inc"] = safe(lambda: t.income_stmt, retries=3, delay=5)
    if d["inc"] is None or d["inc"].empty:
        print(f"  {code}: 无年度利润表"); return None
    time.sleep(1)
    d["q_inc"] = safe(lambda: t.quarterly_income_stmt)
    time.sleep(1)
    d["bs"] = safe(lambda: t.balance_sheet)
    time.sleep(1)
    d["q_bs"] = safe(lambda: t.quarterly_balance_sheet)
    time.sleep(1)
    try:
        fi = t.fast_info
        d["price"], d["mcap_raw"], d["shares"], d["quote_ccy"] = fi.last_price, fi.market_cap, fi.shares, fi.currency
    except Exception:
        d["price"] = d["mcap_raw"] = d["shares"] = None; d["quote_ccy"] = ""

    d["fin_ccy"] = KNOWN_FIN_CCY.get(code, d.get("quote_ccy", ""))
    if not d["fin_ccy"]:
        try:
            d["fin_ccy"] = t.info.get("financialCurrency") or d.get("quote_ccy", "")
        except Exception:
            d["fin_ccy"] = d.get("quote_ccy", "")

    quote_ccy = d.get("quote_ccy", "")
    fin_ccy = d.get("fin_ccy", "")
    if d["mcap_raw"] and quote_ccy and fin_ccy and quote_ccy != fin_ccy:
        fx = get_fx(quote_ccy, fin_ccy)
        d["mcap"] = d["mcap_raw"] * fx
        d["ccy"] = fin_ccy
        print(f"  MCap 换算: {quote_ccy} → {fin_ccy} (×{fx:.4f})")
    else:
        d["mcap"] = d["mcap_raw"]
        d["ccy"] = fin_ccy or quote_ccy

    d["eps_est"] = safe(lambda: t.earnings_estimate)
    time.sleep(0.3)
    d["rev_est"] = safe(lambda: t.revenue_estimate)
    time.sleep(0.3)
    try:
        d["name"] = t.info.get("shortName") or code
    except Exception:
        d["name"] = code
    return d


def fetch_one(code):
    """FMP 优先，整体回退 yfinance。"""
    d = fetch_one_fmp(code)
    if d is not None:
        return d
    print(f"  FMP 获取失败，回退 yfinance: {code}")
    return fetch_one_yf(code)


def _get_shares(bs, col):
    return _first(bs, "Ordinary Shares Number", "Share Issued", col=col)


def _extract_period(inc, bs, col):
    """Extract one period's metrics from income stmt + balance sheet."""
    end = col.date() if hasattr(col, 'date') else col
    rev = _first(inc, "Total Revenue", col=col)
    gp = _first(inc, "Gross Profit", col=col)
    ebit = _first(inc, "Operating Income", "EBIT", col=col)
    ni = _first(inc, "Net Income", col=col)
    eps = _first(inc, "Diluted EPS", "Basic EPS", col=col)
    shr = _get_shares(bs, col) if (bs is not None and col in bs.columns) else None
    gm = _ratio(gp, rev)
    opm = _ratio(ebit, rev)
    npm = _ratio(ni, rev)
    return dict(end_date=end, revenue=rev, gm=gm, opm=opm,
                net_income=ni, npm=npm, eps=eps, shares=shr, n_analysts=None)


def _est_val(est_df, period, field="avg"):
    if est_df is None or est_df.empty or period not in est_df.index:
        return None
    v = est_df.loc[period, field]
    if v is not None and not (isinstance(v, float) and np.isnan(v)):
        return float(v)
    return None


def build_annual(d):
    inc, bs = d["inc"], d.get("bs")
    cols = sorted(inc.columns)
    # Need 4 cols: 3 actual + 1 extra for yoy
    use = cols[-(min(4, len(cols))):]
    rows = []
    prev_rev = None
    for c in use:
        p = _extract_period(inc, bs, c)
        fy = p["end_date"].year
        p["yoy"] = _ratio(p["revenue"], prev_rev) - 1 if _ratio(p["revenue"], prev_rev) else None
        p.update(fy=fy, kind="actual", label=str(fy))
        rows.append(p)
        prev_rev = p["revenue"]
    if len(rows) > 3:
        rows = rows[-3:]

    last = rows[-1]
    fy_m, fy_d, last_fy = last["end_date"].month, last["end_date"].day, last["fy"]
    off = 1 if last["end_date"] < dt.date.today() else 0
    shares = d.get("shares")

    for period in ["0y", "+1y"]:
        fy = last_fy + off; off += 1
        fwd_end = dt.date(fy, fy_m, min(fy_d, 28))
        rev = _est_val(d.get("rev_est"), period)
        e_eps = _est_val(d.get("eps_est"), period)
        ni = e_eps * shares if (e_eps and shares) else None
        na_v = _est_val(d.get("rev_est"), period, "numberOfAnalysts")
        na = int(na_v) if na_v else None
        yoy = _ratio(rev, prev_rev) - 1 if _ratio(rev, prev_rev) else None
        npm = _ratio(ni, rev)
        rows.append(dict(fy=fy, kind="estimate", end_date=fwd_end, label=f"{fy}E",
                         revenue=rev, yoy=yoy, gm=None, opm=None,
                         net_income=ni, npm=npm, eps=e_eps, shares=shares, n_analysts=na))
        if rev: prev_rev = rev
    return rows


def build_quarterly(d):
    qi, qb = d.get("q_inc"), d.get("q_bs")
    if qi is None or qi.empty:
        return []
    cols = sorted(qi.columns)
    act_cols = cols[-4:] if len(cols) >= 4 else cols
    rows = []
    for c in act_cols:
        p = _extract_period(qi, qb, c)
        p.update(kind="actual", label=qlabel(p["end_date"]), yoy=None)
        rows.append(p)

    last_end = rows[-1]["end_date"]
    shares = d.get("shares")
    for i, period in enumerate(["0q", "+1q"]):
        fwd_end = (last_end + pd.DateOffset(months=3 * (i + 1)))
        fwd_end = fwd_end.date() if hasattr(fwd_end, 'date') else fwd_end
        rev = _est_val(d.get("rev_est"), period)
        e_eps = _est_val(d.get("eps_est"), period)
        ni = e_eps * shares if (e_eps and shares) else None
        na_v = _est_val(d.get("rev_est"), period, "numberOfAnalysts")
        na = int(na_v) if na_v else None
        npm = _ratio(ni, rev)
        rows.append(dict(kind="estimate", end_date=fwd_end, label=qlabel(fwd_end),
                         revenue=rev, yoy=None, gm=None, opm=None,
                         net_income=ni, npm=npm, eps=e_eps, shares=shares, n_analysts=na))
    return rows


# ══════════════════════════════════════════════════════════════════
def write_sheet(wb, name, dcode, d, annual, quarterly):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as gl

    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)

    f = Font(name="Arial", size=10)
    fb = Font(name="Arial", size=10, bold=True)
    fbl = Font(name="Arial", size=10, color="0000FF")
    fg = Font(name="Arial", size=9, color="808080")
    hd = PatternFill("solid", start_color="DDEBF7")
    ye = PatternFill("solid", start_color="FFF2CC")
    ctr = Alignment(horizontal="center")
    indent1 = Alignment(indent=1)

    ccy = d.get("ccy", "USD")

    ann_a = [r for r in annual if r["kind"] == "actual"][-3:]
    ann_e = [r for r in annual if r["kind"] == "estimate"][:2]
    ann = ann_a + ann_e
    q_a = [r for r in quarterly if r["kind"] == "actual"][-4:]
    q_e = [r for r in quarterly if r["kind"] == "estimate"][:2]
    qtr = q_a + q_e

    # Row 3: Title
    ws["B3"] = f"{dcode}, {d.get('name', dcode)}"; ws["B3"].font = fb

    # Row 4: End in (estimates only)
    ws["B4"] = "End in"; ws["B4"].font = fb
    for items, start in [(ann, ANN_START), (qtr, Q_START)]:
        for i, r in enumerate(items):
            if r["kind"] == "estimate" and r.get("end_date"):
                c = ws.cell(row=4, column=start + i, value=r["end_date"])
                c.number_format = "yyyy/m/d"; c.font = fb

    # Row 5: Period labels
    ws["B5"] = "Fiscal year"; ws["B5"].font = fb
    for items, start in [(ann, ANN_START), (qtr, Q_START)]:
        for i, r in enumerate(items):
            c = ws.cell(row=R_FY, column=start + i, value=r["label"])
            c.font = fb; c.fill = hd; c.alignment = ctr

    # Data rows
    METRICS = [
        (R_REV, "Revenue",           "revenue",    "#,##0",  True,  False),
        (R_YOY, "yoy",               "yoy",        "0.0%",   False, False),
        (R_GM,  "GM",                "gm",         "0.0%",   False, False),
        (R_OPM, "OPM (EBIT margin)", "opm",        "0.0%",   False, False),
        (R_NI,  "Net Profit (GAAP)", "net_income",  "#,##0", True,  False),
        (R_NPM, "NPM",              "npm",         "0.0%",   False, True),  # indented
        (R_EPS, "EPS (diluted)",     "eps",         "0.00",  False, False),
        (R_SHR, "Shares Out (mm)",   "shares",      "#,##0.0", True,  False),
    ]

    def _wr(ws, row, key, nfmt, is_mm, items, col0):
        for i, r in enumerate(items):
            col = col0 + i
            v = r.get(key)
            if v is not None and not (isinstance(v, float) and np.isnan(v)):
                if is_mm and key == "shares":
                    ws.cell(row=row, column=col, value=round(v / MM, 1))
                elif is_mm and key in ("revenue", "net_income"):
                    ws.cell(row=row, column=col, value=round(v / MM, 0))
                else:
                    ws.cell(row=row, column=col, value=v)
                ws.cell(row=row, column=col).number_format = nfmt
            ws.cell(row=row, column=col).font = fbl if r["kind"] == "actual" else f
            if r["kind"] == "estimate":
                ws.cell(row=row, column=col).fill = ye

    for row, label, key, nfmt, is_mm, indented in METRICS:
        c = ws.cell(row=row, column=2, value=label)
        c.font = fb
        if indented:
            c.alignment = indent1
        _wr(ws, row, key, nfmt, is_mm, ann, ANN_START)
        _wr(ws, row, key, nfmt, is_mm, qtr, Q_START)

    # Valuation block
    ws.cell(row=R_DATE, column=2, value="Date").font = fb
    ws.cell(row=R_DATE, column=3, value=dt.date.today()).number_format = "yyyy-mm-dd"

    q_ccy = d.get("quote_ccy", ccy)
    ws.cell(row=R_PRICE, column=2, value=f"Share price ({q_ccy})").font = fb
    if d.get("price"):
        c = ws.cell(row=R_PRICE, column=3, value=round(d["price"], 2))
        c.font = fbl; c.number_format = "#,##0.00"

    ws.cell(row=R_MCAP, column=2, value=f"Market Cap ({ccy} mm)").font = fb
    if d.get("mcap"):
        c = ws.cell(row=R_MCAP, column=3, value=round(d["mcap"] / MM, 0))
        c.font = fbl; c.number_format = "#,##0"

    # Forward PE / PS formulas
    ws.cell(row=R_PE, column=2, value="Forward PE").font = fb
    ws.cell(row=R_PS, column=2, value="Forward PS").font = fb
    for i, r in enumerate(ann):
        col = ANN_START + i
        cl = gl(col)
        ws.cell(row=R_PE, column=col,
                value=f'=IF(OR({cl}{R_NI}="",{cl}{R_NI}<=0),"",C${R_MCAP}/{cl}{R_NI})')
        ws.cell(row=R_PE, column=col).number_format = '0.0"x"'
        ws.cell(row=R_PS, column=col,
                value=f'=IF(OR({cl}{R_REV}="",{cl}{R_REV}=0),"",C${R_MCAP}/{cl}{R_REV})')
        ws.cell(row=R_PS, column=col).number_format = '0.0"x"'
        if r["kind"] == "estimate":
            ws.cell(row=R_PE, column=col).fill = ye
            ws.cell(row=R_PS, column=col).fill = ye

    # Analysts
    ws.cell(row=R_ANLST, column=2, value="# Analysts (rev)").font = fg
    for items, start in [(ann, ANN_START), (qtr, Q_START)]:
        for i, r in enumerate(items):
            if r.get("n_analysts"):
                ws.cell(row=R_ANLST, column=start + i, value=r["n_analysts"]).font = fg

    # Widths
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 22
    for j in range(ANN_START, ANN_START + len(ann)):
        ws.column_dimensions[gl(j)].width = 12
    ws.column_dimensions["I"].width = 2
    for j in range(Q_START, Q_START + len(qtr)):
        ws.column_dimensions[gl(j)].width = 12


# ══════════════════════════════════════════════════════════════════
def write_summary(wb, succeeded):
    """Write Summary sheet with Forward PE/PS for both sectors."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as gl

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary", 0)  # first sheet

    fb = Font(name="Arial", size=10, bold=True)
    f = Font(name="Arial", size=10)
    hd = PatternFill("solid", start_color="DDEBF7")
    ye = PatternFill("solid", start_color="FFF2CC")

    sw_ok = [d for d in SPORTSWEAR if d[0] in succeeded]
    lux_ok = [d for d in LUXURY if d[0] in succeeded]

    def _write_block(ws, start_row, title, tickers):
        r = start_row
        # PE section
        ws.cell(row=r, column=1, value=title).font = Font(name="Arial", size=11, bold=True)
        r += 1
        ws.cell(row=r, column=1, value="Forward PE").font = fb
        for j, h in enumerate(["FY-2", "FY-1", "FY", "FY+1E", "FY+2E"], 3):
            c = ws.cell(row=r, column=j, value=h)
            c.font = fb; c.fill = hd; c.alignment = Alignment(horizontal="center")
        r += 1
        pe_start = r
        for disp, _ in tickers:
            ws.cell(row=r, column=2, value=disp).font = f
            sn = disp.replace("'", "''")  # escape for formulas
            for j in range(5):
                col_letter = gl(ANN_START + j)
                ws.cell(row=r, column=3 + j,
                        value=f"=IF(ISBLANK('{sn}'!{col_letter}{R_PE}),\"\",'{sn}'!{col_letter}{R_PE})")
                ws.cell(row=r, column=3 + j).number_format = "0.0"
                if j >= 3:
                    ws.cell(row=r, column=3 + j).fill = ye
            r += 1
        r += 1

        # PS section
        ws.cell(row=r, column=1, value="Forward PS").font = fb
        for j, h in enumerate(["FY-2", "FY-1", "FY", "FY+1E", "FY+2E"], 3):
            c = ws.cell(row=r, column=j, value=h)
            c.font = fb; c.fill = hd; c.alignment = Alignment(horizontal="center")
        r += 1
        for disp, _ in tickers:
            ws.cell(row=r, column=2, value=disp).font = f
            sn = disp.replace("'", "''")
            for j in range(5):
                col_letter = gl(ANN_START + j)
                ws.cell(row=r, column=3 + j,
                        value=f"=IF(ISBLANK('{sn}'!{col_letter}{R_PS}),\"\",'{sn}'!{col_letter}{R_PS})")
                ws.cell(row=r, column=3 + j).number_format = "0.0"
                if j >= 3:
                    ws.cell(row=r, column=3 + j).fill = ye
            r += 1
        return r + 2

    next_row = _write_block(ws, 2, "Sportswear", sw_ok)
    _write_block(ws, next_row, "Luxury", lux_ok)

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 14
    for j in range(3, 8):
        ws.column_dimensions[gl(j)].width = 10


# ══════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tickers", nargs="*")
    ap.add_argument("--xlsx", default="output/monitor/Valuation screening.xlsx")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--list", action="store_true")
    a = ap.parse_args()

    if a.list:
        for d, y in ALL_TICKERS: print(f"  {d:14s} → {y}")
        return

    if a.tickers:
        sel = [(d, D2Y[d]) for d in a.tickers if d in D2Y]
        if not sel: sys.exit("未找到 ticker。用 --list 查看。")
    elif a.all:
        sel = ALL_TICKERS
    else:
        sel = ALL_TICKERS[:10]

    from openpyxl import load_workbook, Workbook
    try:
        wb = load_workbook(a.xlsx)
    except FileNotFoundError:
        wb = Workbook()

    ok, fail, succeeded = 0, 0, set()
    total = len(sel)
    for idx, (disp, yahoo) in enumerate(sel):
        print(f"\n--- {disp} ({yahoo}) ---")
        d = fetch_one(yahoo)
        if not d:
            fail += 1; continue
        try:
            ann = build_annual(d)
            qtr = build_quarterly(d)
            write_sheet(wb, disp, disp, d, ann, qtr)
            mc = d.get("mcap", 0)
            print(f"  {d.get('name')} | {d.get('ccy')} | MCap {mc/MM:,.0f}mm")
            a_fy = [r["label"] for r in ann if r["kind"] == "actual"]
            e_fy = [r["label"] for r in ann if r["kind"] == "estimate"]
            q_lab = [r["label"] for r in qtr]
            print(f"  Annual: {a_fy} + {e_fy} | Q: {q_lab}")
            succeeded.add(disp)
            ok += 1
        except Exception as e:
            print(f"  失败: {e}")
            import traceback; traceback.print_exc()
            fail += 1
        # 每个 ticker 之间暂停，避免限速
        if idx < total - 1:
            time.sleep(2)

    # Also include previously existing sheets
    for d, _ in ALL_TICKERS:
        if d in wb.sheetnames:
            succeeded.add(d)

    print(f"\n--- 生成 Summary ---")
    write_summary(wb, succeeded)

    wb.save(a.xlsx)
    print(f"\n完成：{ok} 成功, {fail} 失败 → {a.xlsx}")


if __name__ == "__main__":
    main()
