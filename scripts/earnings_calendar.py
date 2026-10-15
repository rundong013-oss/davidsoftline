#!/usr/bin/env python3
"""
earnings_calendar.py — 运动服饰 & 奢侈品 Earnings 日历

通过 Nasdaq Calendar API 扫描未来 6 个月的 earnings 日期，
匹配覆盖公司，输出到 Excel（运动服饰 / 奢侈品 两个 sheet）。

Usage:
    python3 earnings_calendar.py
    python3 earnings_calendar.py --xlsx path/to/output.xlsx
    python3 earnings_calendar.py --months 9    # 扫描更长时间
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Ticker 清单 ──────────────────────────────────────────────────
# (显示名, Nasdaq symbol, 公司名)
# Nasdaq calendar 只覆盖美股，非美股用卖方/公告日期补充

SPORTSWEAR = [
    ("NKE", "NKE", "Nike"),
    ("ADS.DE", None, "Adidas"),
    ("LULU", "LULU", "Lululemon"),
    ("ONON", "ONON", "On Holding"),
    ("DECK", "DECK", "Deckers/HOKA"),
    ("CROX", "CROX", "Crocs"),
    ("UAA", "UAA", "Under Armour"),
    ("RL", "RL", "Ralph Lauren"),
    ("COLM", "COLM", "Columbia"),
    ("GAP", "GAP", "Gap"),
    ("ANF", "ANF", "Abercrombie"),
    ("PVH", "PVH", "PVH (Calvin Klein)"),
    ("TPR", "TPR", "Tapestry"),
    ("CPRI", "CPRI", "Capri (MK)"),
    ("LEVI", "LEVI", "Levi Strauss"),
    ("KTB", "KTB", "Kontoor"),
    ("GIL", "GIL", "Gildan"),
    ("VFC", "VFC", "VF Corp"),
    ("SHOO", "SHOO", "Steven Madden"),
    ("WWW", "WWW", "Wolverine"),
    ("NWL", "NWL", "Newell"),
    ("ZGN", "ZGN", "Zegna"),
    ("AS", "AS", "Amer Sports"),
    ("PUM.DE", None, "Puma"),
    ("2020.HK", None, "安踏体育"),
    ("2331.HK", None, "李宁"),
    ("3998.HK", None, "波司登"),
    ("1368.HK", None, "特步"),
    ("1361.HK", None, "361度"),
    ("9983.T", None, "迅销(优衣库)"),
    ("002832.SZ", None, "比音勒芬"),
]

LUXURY = [
    ("MC.PA", None, "LVMH"),
    ("RMS.PA", None, "Hermès"),
    ("KER.PA", None, "Kering"),
    ("CFR.SW", None, "Richemont"),
    ("1913.HK", None, "Prada"),
    ("MONC.MI", None, "Moncler"),
    ("SFER.MI", None, "Ferragamo"),
    ("BRBY.L", None, "Burberry"),
    ("RACE", "RACE", "Ferrari"),
    ("GOOS.TO", None, "Canada Goose"),
]


# ── Nasdaq Calendar API ──────────────────────────────────────────
def _get_session():
    """创建 HTTP session。"""
    try:
        from curl_cffi.requests import Session
        return Session(impersonate="chrome")
    except ImportError:
        import requests
        s = requests.Session()
        s.headers.update({
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                          "AppleWebKit/537.36",
            "Accept": "application/json",
        })
        return s


def fetch_calendar_date(session, date: dt.date) -> list[dict]:
    """获取某一天的 Nasdaq earnings calendar。"""
    url = f"https://api.nasdaq.com/api/calendar/earnings?date={date}"
    try:
        r = session.get(url, headers={"Accept": "application/json"},
                        timeout=10)
        if r.status_code != 200:
            return []
        data = r.json()
        rows = data.get("data", {}).get("rows", [])
        return rows or []
    except Exception:
        return []


def scan_earnings(session, months: int = 6) -> dict:
    """
    扫描未来 N 个月的交易日，收集所有 earnings 日期。
    返回 {symbol: [{"date": date, "eps_forecast": str, ...}, ...]}
    """
    today = dt.date.today()
    end = today + dt.timedelta(days=months * 31)

    # 生成所有工作日
    dates = []
    d = today
    while d <= end:
        if d.weekday() < 5:  # 周一到周五
            dates.append(d)
        d += dt.timedelta(days=1)

    print(f"  扫描 {len(dates)} 个交易日 ({today} → {end})...")

    results: dict[str, list] = {}

    def _fetch(date):
        time.sleep(0.15)
        rows = fetch_calendar_date(session, date)
        return date, rows

    done = 0
    with ThreadPoolExecutor(max_workers=6) as pool:
        futures = [pool.submit(_fetch, d) for d in dates]
        for f in as_completed(futures):
            date, rows = f.result()
            done += 1
            for row in rows:
                sym = row.get("symbol", "")
                if sym:
                    if sym not in results:
                        results[sym] = []
                    results[sym].append({
                        "date": date,
                        "name": row.get("name", ""),
                        "eps_forecast": row.get("epsForecast", ""),
                        "fiscal_qtr": row.get("fiscalQuarterEnding", ""),
                        "time": row.get("time", ""),
                    })
            if done % 20 == 0:
                print(f"  ... {done}/{len(dates)} 天已扫描")

    print(f"  扫描完成，共发现 {len(results)} 家公司有 earnings")
    return results


def match_tickers(calendar: dict, ticker_list: list[tuple]) -> dict:
    """将 calendar 数据匹配到我们的 ticker 清单。"""
    matched = {}
    for display, nasdaq_sym, name in ticker_list:
        earnings = []
        if nasdaq_sym and nasdaq_sym in calendar:
            for e in sorted(calendar[nasdaq_sym],
                            key=lambda x: x["date"]):
                earnings.append({
                    "date": dt.datetime.combine(
                        e["date"], dt.time(0, 0),
                        tzinfo=dt.timezone.utc),
                    "eps_estimate": e["eps_forecast"],
                    "fiscal_qtr": e["fiscal_qtr"],
                    "time": e["time"],
                })
        matched[display] = {
            "name": name,
            "nasdaq_sym": nasdaq_sym,
            "earnings": earnings[:2],  # 只保留最近 2 次
        }
    return matched


# ── Excel 输出 ───────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
DATE_SOON_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
HEADER_FONT = Font(name="Arial", size=10, bold=True)
BODY_FONT = Font(name="Arial", size=10)
GRAY_FONT = Font(name="Arial", size=9, color="999999")


def write_sheet(ws, data: dict):
    """将一组公司的 earnings 数据写入 worksheet。"""
    headers = ["Ticker", "公司",
               "下次 Earnings", "距今天数", "EPS Est", "报告季度",
               "再下次 Earnings", "距今天数", "EPS Est", "报告季度",
               "更新时间"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")

    now = dt.datetime.now(tz=dt.timezone.utc)
    update_str = now.strftime("%Y-%m-%d")

    def sort_key(item):
        e = item[1]["earnings"]
        if e and hasattr(e[0]["date"], "timestamp"):
            return (0, e[0]["date"].timestamp())
        return (1, 0)

    sorted_items = sorted(data.items(), key=sort_key)

    for row_idx, (ticker, info) in enumerate(sorted_items, 2):
        ws.cell(row=row_idx, column=1, value=ticker).font = Font(
            name="Arial", size=10, bold=True)
        ws.cell(row=row_idx, column=2, value=info["name"]).font = BODY_FONT

        for i, (d_col, day_col, eps_col, qtr_col) in enumerate(
                [(3, 4, 5, 6), (7, 8, 9, 10)]):
            if i < len(info["earnings"]):
                e = info["earnings"][i]
                date_val = e["date"]
                if hasattr(date_val, "strftime"):
                    c = ws.cell(row=row_idx, column=d_col,
                                value=date_val.replace(tzinfo=None))
                    c.number_format = "YYYY-MM-DD"
                    days = (date_val - now).days
                    c_days = ws.cell(row=row_idx, column=day_col, value=days)
                    if 0 <= days <= 30:
                        c.fill = DATE_SOON_FILL
                        c_days.fill = DATE_SOON_FILL
                eps = e.get("eps_estimate", "")
                if eps:
                    ws.cell(row=row_idx, column=eps_col, value=eps)
                qtr = e.get("fiscal_qtr", "")
                if qtr:
                    ws.cell(row=row_idx, column=qtr_col, value=qtr)

        # 无 Nasdaq 覆盖的标注
        if not info.get("nasdaq_sym"):
            ws.cell(row=row_idx, column=3, value="非美股").font = GRAY_FONT

        ws.cell(row=row_idx, column=11, value=update_str).font = GRAY_FONT

        for col in range(1, 12):
            cell = ws.cell(row=row_idx, column=col)
            if cell.font is None or (not cell.font.bold and cell.font.color
                                     is None):
                cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    widths = [12, 18, 16, 10, 10, 14, 16, 10, 10, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_excel(sw_data: dict, lx_data: dict, xlsx_path: str):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "运动服饰"
    write_sheet(ws1, sw_data)

    ws2 = wb.create_sheet("奢侈品")
    write_sheet(ws2, lx_data)

    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"\n✓ 已保存: {xlsx_path}")


# ── Main ─────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(description="Earnings 日历生成器")
    ap.add_argument("--xlsx",
                    default="output/monitor/Earnings calendar.xlsx")
    ap.add_argument("--months", type=int, default=6,
                    help="扫描未来几个月（默认 6）")
    args = ap.parse_args()

    session = _get_session()

    print("=== 扫描 Nasdaq Earnings Calendar ===")
    calendar = scan_earnings(session, months=args.months)

    print("\n=== 匹配运动服饰 ===")
    sw_data = match_tickers(calendar, SPORTSWEAR)
    for ticker, info in sw_data.items():
        n = len(info["earnings"])
        status = "✓" if n else ("—" if info["nasdaq_sym"] else "·")
        dates = ", ".join(e["date"].strftime("%Y-%m-%d")
                         for e in info["earnings"])
        extra = f"  → {dates}" if dates else ""
        print(f"  {status} {ticker:12s} {info['name']}{extra}")

    print("\n=== 匹配奢侈品 ===")
    lx_data = match_tickers(calendar, LUXURY)
    for ticker, info in lx_data.items():
        n = len(info["earnings"])
        status = "✓" if n else ("—" if info["nasdaq_sym"] else "·")
        dates = ", ".join(e["date"].strftime("%Y-%m-%d")
                         for e in info["earnings"])
        extra = f"  → {dates}" if dates else ""
        print(f"  {status} {ticker:12s} {info['name']}{extra}")

    build_excel(sw_data, lx_data, args.xlsx)

    # 未来 30 天摘要
    now = dt.datetime.now(tz=dt.timezone.utc)
    upcoming = []
    for group, label in [(sw_data, "运动服饰"), (lx_data, "奢侈品")]:
        for ticker, info in group.items():
            for e in info["earnings"]:
                days = (e["date"] - now).days
                if 0 <= days <= 30:
                    upcoming.append((e["date"], ticker, info["name"],
                                     days, label))
    if upcoming:
        upcoming.sort()
        print("\n=== 未来 30 天内 Earnings ===")
        for d, t, n, days, lbl in upcoming:
            print(f"  {d.strftime('%Y-%m-%d')}  {t:12s} {n:20s} "
                  f"({days}天后) [{lbl}]")

    # 覆盖统计
    sw_ok = sum(1 for v in sw_data.values() if v["earnings"])
    sw_us = sum(1 for _, s, _ in SPORTSWEAR if s)
    lx_ok = sum(1 for v in lx_data.values() if v["earnings"])
    lx_us = sum(1 for _, s, _ in LUXURY if s)
    print(f"\n覆盖率: 运动服饰 {sw_ok}/{sw_us} 美股, "
          f"奢侈品 {lx_ok}/{lx_us} 美股")
    non_us = sum(1 for _, s, _ in SPORTSWEAR + LUXURY if not s)
    print(f"        非美股 {non_us} 家（需手动补充）")


if __name__ == "__main__":
    main()
